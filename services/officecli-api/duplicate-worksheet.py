import copy
import json
import os
import re
import sys
import tempfile

from openpyxl import load_workbook


INVALID_WORKSHEET_CHARACTER = re.compile(r"[\\/*?:\[\]]")


def worksheet_name(value, field_name):
    name = value.strip()
    if not name:
        raise ValueError(f"{field_name} is required")
    if len(name) > 31:
        raise ValueError(f"{field_name} must be at most 31 characters")
    if INVALID_WORKSHEET_CHARACTER.search(name):
        raise ValueError(f"{field_name} contains a character Excel does not allow")
    return name


def find_worksheet(workbook, name):
    folded = name.casefold()
    return next((sheet for sheet in workbook.worksheets if sheet.title.casefold() == folded), None)


def assert_supported_content(source):
    unsupported = []
    if source._images:
        unsupported.append("images")
    if source._charts:
        unsupported.append("charts")
    if source.tables:
        unsupported.append("tables")
    if source._pivots:
        unsupported.append("pivot tables")
    if unsupported:
        raise ValueError(
            "The source worksheet contains content that cannot be cloned safely: " + ", ".join(unsupported)
        )


def copy_extended_settings(source, target):
    target.auto_filter = copy.copy(source.auto_filter)
    target.conditional_formatting = copy.deepcopy(source.conditional_formatting)
    target.data_validations = copy.deepcopy(source.data_validations)
    target.protection = copy.copy(source.protection)
    target.views = copy.copy(source.views)
    target.freeze_panes = source.freeze_panes
    target.row_breaks = copy.deepcopy(source.row_breaks)
    target.col_breaks = copy.deepcopy(source.col_breaks)
    target.sheet_state = source.sheet_state
    target.print_area = source.print_area
    target.print_title_cols = source.print_title_cols
    target.print_title_rows = source.print_title_rows


def duplicate_worksheet(document, source_name, destination_name):
    source_name = worksheet_name(source_name, "sourceWorksheet")
    destination_name = worksheet_name(destination_name, "destinationWorksheet")
    workbook = load_workbook(document, data_only=False, keep_links=True)
    source = find_worksheet(workbook, source_name)
    if source is None:
        raise ValueError(f"Worksheet not found: {source_name}")
    if find_worksheet(workbook, destination_name) is not None:
        raise ValueError(f"Worksheet already exists: {destination_name}")

    assert_supported_content(source)
    target = workbook.copy_worksheet(source)
    target.title = destination_name
    copy_extended_settings(source, target)

    workbook._sheets.remove(target)
    workbook._sheets.insert(workbook._sheets.index(source) + 1, target)
    position = workbook._sheets.index(target)

    file_descriptor, temporary_path = tempfile.mkstemp(
        prefix=".officecli-duplicate-", suffix=".xlsx", dir=os.path.dirname(document)
    )
    os.close(file_descriptor)
    try:
        workbook.save(temporary_path)
        os.replace(temporary_path, document)
    finally:
        workbook.close()
        if os.path.exists(temporary_path):
            os.unlink(temporary_path)

    return {
        "document": os.path.basename(document),
        "sourceWorksheet": source_name,
        "destinationWorksheet": destination_name,
        "position": position,
    }


def main():
    if len(sys.argv) != 4:
        raise ValueError("Expected document, source worksheet, and destination worksheet arguments")
    result = duplicate_worksheet(sys.argv[1], sys.argv[2], sys.argv[3])
    print(json.dumps({"success": True, "data": result}, separators=(",", ":")))


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(json.dumps({"success": False, "error": {"message": str(error)}}, separators=(",", ":")))
        sys.exit(1)
